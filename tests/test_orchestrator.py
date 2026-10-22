"""Error recovery inside the agent pipeline."""

import asyncio
import json

import pytest
from docx import Document

from app.llm import LLMClient, LLMUnavailableError
from app.llm.mock import MockProvider
from app.orchestrator import DocumentAgent
from tests.conftest import STANDARD_REQUEST, ScriptedProvider, make_settings


class FailingWriterProvider(MockProvider):
    """Behaves like the mock, except the writer always fails for one section."""

    name = "mock-failing-writer"

    async def chat(self, messages, **kwargs):
        if "AGENT: WRITER" in messages[0]["content"] and "Budget" in messages[1]["content"][:80]:
            raise LLMUnavailableError("simulated outage")
        return await super().chat(messages, **kwargs)


class ScoreDropReviewerProvider(MockProvider):
    """Reviewer gives 7 first, then 5 after revision -> agent must keep the original."""

    name = "mock-score-drop"

    async def chat(self, messages, **kwargs):
        reply = await super().chat(messages, **kwargs)
        if "AGENT: REVIEWER" in messages[0]["content"] and '"review_round": 2' in messages[1]["content"]:
            return json.dumps({"score": 5, "summary": "Worse.", "issues": []})
        return reply


class CrossSectionIssueProvider(MockProvider):
    """Round 1 flags a date mismatch in section 1 that also affects section 3."""

    name = "mock-cross-section"

    async def chat(self, messages, **kwargs):
        if "AGENT: REVIEWER" in messages[0]["content"] and '"review_round": 1' in messages[1]["content"]:
            return json.dumps({"score": 6, "summary": "Dates disagree.", "issues": [{
                "section_id": 1, "related_section_ids": [3, 99, 1], "severity": "major",
                "problem": "Completion year differs between sections 1 and 3.",
                "fix_instruction": "Use 2027-02-28 everywhere."}]})
        return await super().chat(messages, **kwargs)


def test_cross_section_issue_revises_every_related_section(tmp_path):
    agent = make_agent(tmp_path, CrossSectionIssueProvider())
    result = asyncio.run(agent.run(STANDARD_REQUEST))
    revise_tasks = [t.title for t in result.todo if t.title.startswith("Revise section")]
    assert [t.split(":")[0] for t in revise_tasks] == ["Revise section 1", "Revise section 3"]
    assert result.review.final_score == 9


class PassingScoreWithMajorIssueProvider(MockProvider):
    """Round 1 scores 9 (passes) but still reports a contradiction."""

    name = "mock-major-issue"

    async def chat(self, messages, **kwargs):
        if "AGENT: REVIEWER" in messages[0]["content"] and '"review_round": 1' in messages[1]["content"]:
            return json.dumps({"score": 9, "summary": "Good, one contradiction.", "issues": [{
                "section_id": 2, "severity": "major",
                "problem": "Says 4 hotel nights; the rest of the plan says 7.",
                "fix_instruction": "Use 7 nights."}]})
        return await super().chat(messages, **kwargs)


def test_major_issue_is_fixed_even_when_score_passes(tmp_path):
    agent = make_agent(tmp_path, PassingScoreWithMajorIssueProvider())
    result = asyncio.run(agent.run(STANDARD_REQUEST))
    assert any(t.title.startswith("Revise section 2") for t in result.todo)
    assert result.review.revised_sections


def make_agent(tmp_path, provider) -> DocumentAgent:
    settings = make_settings(tmp_path)
    return DocumentAgent(settings, LLMClient([provider], max_retries=0, base_delay=0))


def test_failed_section_gets_placeholder_and_run_still_completes(tmp_path):
    agent = make_agent(tmp_path, FailingWriterProvider())
    result = asyncio.run(agent.run(STANDARD_REQUEST))

    assert result.status == "completed_with_warnings"
    failed = [t for t in result.todo if t.status == "failed"]
    assert len(failed) == 1 and "Budget" in failed[0].title
    doc = Document(str(tmp_path / result.document.filename))
    assert any("[Draft pending]" in p.text for p in doc.paragraphs)


def test_revision_that_lowers_score_is_discarded(tmp_path):
    agent = make_agent(tmp_path, ScoreDropReviewerProvider())
    result = asyncio.run(agent.run(STANDARD_REQUEST))

    assert result.review.initial_score == 7
    assert result.review.final_score == 7
    assert result.review.revised_sections == []
    assert any("previous version was kept" in n for n in result.review.notes)
    # 7 is still below the 8 target, so the open issue is reported, not hidden.
    assert result.status == "completed_with_warnings"
    assert result.review.remaining_issues and "owner" in result.review.remaining_issues[0]
    doc = Document(str(tmp_path / result.document.filename))
    assert any("Open points the reviewer flagged" in p.text for p in doc.paragraphs)


def test_planner_failure_propagates(tmp_path):
    agent = make_agent(tmp_path, ScriptedProvider("dead", ["nope"]))
    with pytest.raises(LLMUnavailableError):
        asyncio.run(agent.run(STANDARD_REQUEST))
